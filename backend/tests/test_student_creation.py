"""Provisioning regression tests using real ORM transactions and HTTP routes.

Every test creates its own temporary SQLite database; neither the configured
DATABASE_URL nor any real email service is used. SQLite-specific timestamp
and JSONB adaptations below are test-only. PostgreSQL migrations are checked
separately; SQLite does not prove PostgreSQL locking behavior.
"""
import importlib.util
import io
import tempfile
import unittest
import uuid
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.dialects import postgresql
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateIndex

from app.core import student_credentials
from app.core.email import SendResult
from app.core.rate_limit import limiter
from app.core.security import create_access_token, hash_password, verify_password
from app.database import Base, get_db
from app.main import app
from app.models import (
    AdminRole, AdminRolePermission, AuthAuditLog, AuthSession, EnrollmentStatusEnum,
    Permission, Program, ProgramEnrollment, ProgramEnum, ProgramStatusEnum,
    RoleEnum, StudentProfile, SystemSettings, User,
)
from app.schemas.admin_students import CreateStudentRequest


@compiles(JSONB, "sqlite")
def _sqlite_jsonb(type_, compiler, **kwargs):
    return "JSON"


class ProvisioningTestSession(Session):
    pass


@event.listens_for(ProvisioningTestSession, "loaded_as_persistent")
def _sqlite_utc_dates(session, instance):
    # SQLite drops tzinfo. The production PostgreSQL driver preserves it.
    if session.get_bind().dialect.name == "sqlite":
        for key, value in list(vars(instance).items()):
            if isinstance(value, datetime) and value.tzinfo is None:
                setattr(instance, key, value.replace(tzinfo=timezone.utc))


PASSWORD = "Arun@12345"
ADMIN_PASSWORD = "Admin@Test123"


class StudentCreationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="arckenites-student-tests-")
        self.engine = create_engine(f"sqlite:///{Path(self.temp.name) / 'portal.sqlite'}", connect_args={"check_same_thread": False})
        @event.listens_for(self.engine, "connect")
        def _foreign_keys(connection, record):
            connection.execute("PRAGMA foreign_keys=ON")
        Base.metadata.create_all(self.engine)
        with ProvisioningTestSession(self.engine) as db:
            db.add(AdminRole(id=1, name="Super Admin", slug="super_admin", is_system=True))
            db.flush()
            admin = User(username="admin.existing", full_name="Existing Admin", role=RoleEnum.admin,
                         admin_role_id=1, password_hash=hash_password(ADMIN_PASSWORD), must_change_password=False)
            trainer = User(username="trainer.existing", full_name="Existing Trainer", role=RoleEnum.staff,
                           password_hash=hash_password(ADMIN_PASSWORD), must_change_password=False)
            old = User(username="legacy.student", full_name="Legacy Student", role=RoleEnum.student,
                       password_hash=hash_password(PASSWORD), must_change_password=False)
            db.add_all([admin, trainer, old, SystemSettings(id=1)])
            db.flush()
            db.add(StudentProfile(user_id=old.id, student_code="AK-STU-0001", email="legacy@example.com",
                                  program=ProgramEnum.institutional, phone="9876543210"))
            program = Program(slug="python_full_stack", name="Python Full Stack", status=ProgramStatusEnum.active)
            institutional = Program(slug="institutional", name="Institutional Program", status=ProgramStatusEnum.active)
            inactive = Program(slug="inactive_program", name="Inactive", status=ProgramStatusEnum.inactive)
            db.add_all([program, institutional, inactive])
            db.commit()
            self.admin_id, self.trainer_id, self.legacy_id = admin.id, trainer.id, old.id
            self.program_id, self.institutional_id, self.inactive_id = program.id, institutional.id, inactive.id
        self.original_overrides = app.dependency_overrides.copy()
        app.dependency_overrides[get_db] = self._db
        self.client = TestClient(app)  # No lifespan: do not start production background workers.
        self.limit_patch = patch.object(limiter, "enabled", False)
        self.limit_patch.start()
        self.send_patch = patch.object(student_credentials, "send_email_detailed", return_value=SendResult(accepted=True))
        self.send = self.send_patch.start()
        self.headers = self._headers(self.admin_id, "admin")

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        app.dependency_overrides.update(self.original_overrides)
        self.send_patch.stop()
        self.limit_patch.stop()
        self.engine.dispose()
        self.temp.cleanup()

    def _db(self):
        with ProvisioningTestSession(self.engine, autoflush=False) as db:
            yield db

    def _headers(self, user_id, role):
        jti = uuid.uuid4()
        token, expiry = create_access_token(user_id, role, jti)
        with ProvisioningTestSession(self.engine) as db:
            db.add(AuthSession(id=jti, user_id=user_id, expires_at=expiry))
            db.commit()
        return {"Authorization": f"Bearer {token}"}

    def _payload(self, **overrides):
        return {"full_name": "  Arun Kumar  ", "email": " Arun@Example.COM ",
                "temp_password": PASSWORD, "program_id": self.program_id, **overrides}

    def _create(self, **overrides):
        return self.client.post("/api/admin/students", json=self._payload(**overrides), headers=self.headers)

    def test_create_normalizes_email_hashes_password_and_enrolls(self):
        response = self._create()
        self.assertEqual(response.status_code, 201, response.text)
        data = response.json()
        self.assertEqual(data["username"], "arun@example.com")
        self.assertEqual(data["email"], "arun@example.com")
        self.assertEqual(data["full_name"], "Arun Kumar")
        self.assertEqual(data["program_name"], "Python Full Stack")
        self.assertEqual(data["program_id"], self.program_id)
        self.assertTrue(data["is_active"])
        self.assertTrue(data["must_change_password"])
        self.assertTrue(data["credential_email_sent"])
        self.assertNotIn(PASSWORD, response.text)
        self.assertNotIn("password_hash", response.text)
        with ProvisioningTestSession(self.engine) as db:
            user = db.get(User, data["id"])
            self.assertEqual(user.role, RoleEnum.student)
            self.assertNotEqual(user.password_hash, PASSWORD)
            self.assertTrue(verify_password(PASSWORD, user.password_hash))
            self.assertIsNone(user.student_profile.program)  # Custom catalogue slug, not arbitrary enum text.
            enrollment = db.scalar(select(ProgramEnrollment).where(ProgramEnrollment.student_id == user.id))
            self.assertEqual(enrollment.program_id, self.program_id)
            self.assertEqual(enrollment.status, EnrollmentStatusEnum.active)
            self.assertEqual(user.student_profile.program_id, self.program_id)
            audit = db.scalar(select(AuthAuditLog).where(AuthAuditLog.target == "arun@example.com"))
            self.assertNotIn(PASSWORD, audit.detail)

    def test_email_observes_committed_user_profile_enrollment_and_audit(self):
        def send_after_commit(to, subject, body, html_body):
            # This separate connection cannot see another transaction's uncommitted rows.
            with ProvisioningTestSession(self.engine) as observer:
                user = observer.scalar(select(User).where(User.username == to))
                self.assertIsNotNone(user)
                self.assertIsNotNone(user.student_profile)
                self.assertEqual(len(user.program_enrollments), 1)
                self.assertIsNotNone(observer.scalar(select(AuthAuditLog).where(AuthAuditLog.target == to)))
            self.assertIn(PASSWORD, body)
            self.assertIn("Python Full Stack", html_body)
            self.assertIn("/login.html", html_body)
            return SendResult(accepted=True)
        self.send.side_effect = send_after_commit
        response = self._create()
        self.assertEqual(response.status_code, 201)
        self.assertTrue(response.json()["credential_email_sent"])
        self.send.assert_called_once()

    def test_duplicate_email_case_and_spaces_are_rejected(self):
        self.assertEqual(self._create().status_code, 201)
        self.send.reset_mock()
        response = self._create(email=" ARUN@example.com ")
        self.assertEqual(response.status_code, 409)
        self.assertIn("email already exists", response.json()["detail"])
        self.send.assert_not_called()

    def test_legacy_student_contact_email_cannot_create_second_account(self):
        response = self._create(email=" LEGACY@example.COM ")
        self.assertEqual(response.status_code, 409)
        self.send.assert_not_called()

    def test_other_role_login_cannot_be_reused(self):
        with ProvisioningTestSession(self.engine) as db:
            db.get(User, self.trainer_id).username = "Existing.Trainer@Example.com"
            db.commit()
        self.assertEqual(self._create(email="existing.trainer@example.com").status_code, 409)
        self.send.assert_not_called()

    def test_missing_and_invalid_fields_do_not_echo_password_or_send_email(self):
        for changes, message in [({"email": "bad-email"}, "valid email"), ({"email": "a@-bad.com"}, "valid email"),
                                 ({"email": "a..b@example.com"}, "valid email"), ({"full_name": "  "}, "Full name"),
                                 ({"temp_password": ""}, "Temporary password"), ({"program_id": None}, "select a program"),
                                 ({"temp_password": "é" * 40}, "72 UTF-8 bytes")]:
            with self.subTest(changes=list(changes)):
                response = self._create(**changes)
                self.assertEqual(response.status_code, 422)
                self.assertIn(message, response.json()["detail"])
                self.assertNotIn(PASSWORD, response.text)
                self.assertNotIn('"input"', response.text)
        self.send.assert_not_called()

    def test_unknown_and_inactive_programs_are_rejected(self):
        for program_id in (999999, self.inactive_id):
            self.assertEqual(self._create(program_id=program_id).status_code, 400)
        self.send.assert_not_called()

    def test_login_normalizes_student_email_and_wrong_password_is_rejected(self):
        self.assertEqual(self._create().status_code, 201)
        response = self.client.post("/api/auth/login", json={"username": " ARUN@EXAMPLE.COM ", "password": PASSWORD})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["user"]["program_name"], "Python Full Stack")
        auth = {"Authorization": f"Bearer {response.json()['access_token']}"}
        me = self.client.get("/api/auth/me", headers=auth)
        self.assertEqual(me.status_code, 200, me.text)
        self.assertEqual(me.json()["username"], "arun@example.com")
        wrong = self.client.post("/api/auth/login", json={"username": "arun@example.com", "password": "Wrong@Test123"})
        self.assertEqual(wrong.status_code, 401)

    def test_existing_admin_trainer_and_legacy_student_can_log_in(self):
        for username, password, role in (("admin.existing", ADMIN_PASSWORD, "admin"),
                                         ("trainer.existing", ADMIN_PASSWORD, "staff"),
                                         ("legacy.student", PASSWORD, "student")):
            response = self.client.post("/api/auth/login", json={"username": username, "password": password})
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()["user"]["role"], role)
        # New normalization is scoped to students, preserving exact admin/trainer behavior.
        for username in ("ADMIN.EXISTING", " TRAINER.EXISTING "):
            self.assertEqual(self.client.post("/api/auth/login", json={"username": username, "password": ADMIN_PASSWORD}).status_code, 401)

    def test_database_failure_rolls_back_all_records_and_does_not_send(self):
        with patch.object(ProvisioningTestSession, "commit", side_effect=RuntimeError("simulated failure")):
            response = self._create()
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json()["detail"], "Unable to create student. Please try again.")
        self.send.assert_not_called()
        with ProvisioningTestSession(self.engine) as db:
            self.assertIsNone(db.scalar(select(User).where(User.username == "arun@example.com")))
            self.assertEqual(db.scalar(select(func.count()).select_from(ProgramEnrollment)), 0)
            self.assertEqual(db.scalar(select(func.count()).select_from(AuthAuditLog)), 0)

    def test_audit_failure_rolls_back_creation_and_does_not_send(self):
        with patch("app.api.routes_admin_students.write_audit_event", side_effect=RuntimeError("audit unavailable")):
            self.assertEqual(self._create().status_code, 500)
        self.send.assert_not_called()
        with ProvisioningTestSession(self.engine) as db:
            self.assertIsNone(db.scalar(select(User).where(User.username == "arun@example.com")))

    def test_email_failure_keeps_student_and_reports_warning(self):
        self.send.return_value = SendResult(accepted=False, transport_unavailable=True)
        response = self._create()
        self.assertEqual(response.status_code, 201)
        self.assertFalse(response.json()["credential_email_sent"])
        self.assertIn("could not be sent", response.json()["detail"])
        with ProvisioningTestSession(self.engine) as db:
            user = db.get(User, response.json()["id"])
            self.assertTrue(verify_password(PASSWORD, user.password_hash))
            self.assertEqual(len(user.program_enrollments), 1)

    def test_unexpected_email_error_does_not_leak_password_in_logs(self):
        self.send.side_effect = RuntimeError(f"provider echoed {PASSWORD}")
        with self.assertLogs("student_credentials", level="WARNING") as captured:
            response = self._create()
        self.assertEqual(response.status_code, 201)
        self.assertFalse(response.json()["credential_email_sent"])
        self.assertNotIn(PASSWORD, " ".join(captured.output))

    def test_create_and_resend_require_authorized_admin(self):
        created = self._create().json()
        for headers, expected in (({}, 401), (self._headers(self.trainer_id, "staff"), 403),
                                  (self._headers(self.legacy_id, "student"), 403)):
            self.assertEqual(self.client.post("/api/admin/students", json=self._payload(), headers=headers).status_code, expected)
            self.assertEqual(self.client.post(f"/api/admin/students/{created['id']}/credentials-email",
                                             json={"temp_password": PASSWORD}, headers=headers).status_code, expected)
        with ProvisioningTestSession(self.engine) as db:
            db.add(AdminRole(id=2, name="Viewer", slug="test_viewer"))
            db.flush()
            db.get(User, self.admin_id).admin_role_id = 2
            db.commit()
        self.assertEqual(self._create(email="other@example.com").status_code, 403)

    def test_resend_verifies_original_password_and_does_not_change_it(self):
        created = self._create().json()
        with ProvisioningTestSession(self.engine) as db:
            original_hash = db.get(User, created["id"]).password_hash
        path = f"/api/admin/students/{created['id']}/credentials-email"
        self.send.reset_mock()
        wrong = self.client.post(path, json={"temp_password": "Wrong@Test123"}, headers=self.headers)
        self.assertEqual(wrong.status_code, 400)
        self.send.assert_not_called()
        sent = self.client.post(path, json={"temp_password": PASSWORD}, headers=self.headers)
        self.assertEqual(sent.status_code, 200)
        self.assertTrue(sent.json()["credential_email_sent"])
        self.assertNotIn(PASSWORD, sent.text)
        self.send.assert_called_once()
        with ProvisioningTestSession(self.engine) as db:
            self.assertEqual(db.get(User, created["id"]).password_hash, original_hash)

    def test_resend_after_password_change_is_blocked(self):
        created = self._create().json()
        with ProvisioningTestSession(self.engine) as db:
            db.get(User, created["id"]).must_change_password = False
            db.commit()
        self.send.reset_mock()
        response = self.client.post(f"/api/admin/students/{created['id']}/credentials-email",
                                    json={"temp_password": PASSWORD}, headers=self.headers)
        self.assertEqual(response.status_code, 400)
        self.send.assert_not_called()

    def test_first_login_password_change_and_profile_completion_still_work(self):
        created = self._create(program_id=self.institutional_id).json()
        auth = self._headers(created["id"], "student")
        response = self.client.post("/api/auth/change-password", headers=auth,
                                    json={"current_password": PASSWORD, "new_password": "Changed@Test123"})
        self.assertEqual(response.status_code, 204)
        profile = self.client.post("/api/students/me/profile", headers=auth,
                                   json={"full_name": "Arun Kumar", "email": "arun@example.com", "mobile_number": "9876543210"})
        self.assertEqual(profile.status_code, 200, profile.text)
        self.assertTrue(profile.json()["profile_completed"])
        self.assertFalse(profile.json()["must_change_password"])
        self.assertEqual(profile.json()["program"], "institutional")
        self.assertEqual(profile.json()["program_name"], "Institutional Program")

    def test_program_edit_uses_enrollment_and_keeps_account_and_sessions(self):
        created = self._create().json()
        auth = self._headers(created["id"], "student")
        response = self.client.patch(f"/api/admin/students/{created['id']}", headers=self.headers,
                                     json={"program_id": self.institutional_id})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["program_name"], "Institutional Program")
        self.assertEqual(response.json()["program"], "institutional")
        self.assertEqual(self.client.get("/api/auth/me", headers=auth).status_code, 200)
        with ProvisioningTestSession(self.engine) as db:
            user = db.get(User, created["id"])
            self.assertEqual(len(user.program_enrollments), 2)
            self.assertEqual(user.username, created["username"])
            self.assertTrue(verify_password(PASSWORD, user.password_hash))

    def test_excel_export_contains_email_and_real_program_preserving_existing_fields(self):
        self._create()
        response = self.client.get("/api/admin/students/export", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        rows = list(book.active.values)
        headers = rows[0]
        row = next(r for r in rows[1:] if r[headers.index("Email")] == "arun@example.com")
        self.assertEqual(row[headers.index("Full Name")], "Arun Kumar")
        self.assertEqual(row[headers.index("Program")], "Python Full Stack")
        self.assertEqual(row[headers.index("Status")], "Active")
        self.assertIn("Username", headers)
        self.assertIn("Created At", headers)

    def test_program_options_support_provisioner_without_programs_view(self):
        with ProvisioningTestSession(self.engine) as db:
            db.add(AdminRole(id=2, name="Provisioner", slug="test_provisioner"))
            db.add(Permission(key="students.create", module="students", action="create"))
            db.flush()
            db.add(AdminRolePermission(admin_role_id=2, permission_key="students.create"))
            db.get(User, self.admin_id).admin_role_id = 2
            db.commit()
        response = self.client.get("/api/admin/students/program-options", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual({row["id"] for row in response.json()}, {self.program_id, self.institutional_id})
        self.assertEqual(self._create().status_code, 201)

    def test_excel_user_text_is_not_a_formula(self):
        self._create(full_name='=HYPERLINK("https://example.com", "Name")')
        response = self.client.get("/api/admin/students/export", headers=self.headers)
        sheet = load_workbook(io.BytesIO(response.content)).active
        name_cells = [row[1] for row in sheet.iter_rows(min_row=2)]
        value = next(cell for cell in name_cells if cell.value.startswith("=HYPERLINK"))
        self.assertEqual(value.data_type, "s")

    def test_unique_index_rejects_concurrent_style_case_duplicate(self):
        self._create()
        with ProvisioningTestSession(self.engine) as db:
            db.add(User(username="ARUN@EXAMPLE.COM", full_name="Duplicate", role=RoleEnum.student,
                        password_hash=hash_password(PASSWORD)))
            with self.assertRaises(IntegrityError):
                db.commit()
            db.rollback()

    def test_html_email_escapes_name_password_and_program(self):
        self._create(full_name="Arun <img src=x onerror=alert(1)>", temp_password="A<&Password1")
        _, _, body, html = self.send.call_args.args
        self.assertIn("A<&Password1", body)
        self.assertIn("A&lt;&amp;Password1", html)
        self.assertNotIn("<img src=x", html)


class StudentEmailMigrationTests(unittest.TestCase):
    def _migration(self):
        path = Path(__file__).parents[1] / "alembic/versions/c6e8a2b4d9f1_student_email_login.py"
        spec = importlib.util.spec_from_file_location("student_email_migration", path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_migration_blocks_duplicates_before_schema_changes(self):
        migration = self._migration()
        with patch.object(migration, "op") as op:
            op.get_bind.return_value.execute.return_value.scalar_one.return_value = 2
            with self.assertRaisesRegex(RuntimeError, "normalized login duplicates"):
                migration.upgrade()
            op.alter_column.assert_not_called()
            op.create_index.assert_not_called()
            op.add_column.assert_not_called()

    def test_clean_migration_only_widens_adds_index_and_nullable_program_reference(self):
        migration = self._migration()
        with patch.object(migration, "op") as op:
            op.get_bind.return_value.execute.return_value.scalar_one.return_value = 0
            migration.upgrade()
            self.assertEqual(op.alter_column.call_args.kwargs["type_"].length, 255)
            self.assertTrue(op.create_index.call_args.kwargs["unique"])
            self.assertTrue(op.add_column.call_args.args[1].nullable)
            op.drop_column.assert_not_called()

    def test_postgres_index_matches_case_normalized_email_constraint(self):
        index = next(i for i in User.__table__.indexes if i.name == "uq_users_email_login_normalized")
        sql = str(CreateIndex(index).compile(dialect=postgresql.dialect()))
        self.assertIn("UNIQUE INDEX", sql)
        self.assertIn("lower(trim(username))", sql)
        self.assertIn("WHERE", sql)

    def test_downgrade_does_not_truncate_long_logins(self):
        migration = self._migration()
        with patch.object(migration, "op") as op:
            op.get_bind.return_value.execute.return_value.scalar_one.return_value = 1
            with self.assertRaisesRegex(RuntimeError, "exceed 120"):
                migration.downgrade()
            op.drop_column.assert_not_called()


if __name__ == "__main__":
    unittest.main()
